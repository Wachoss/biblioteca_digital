"""Fase 1: leer el Excel y normalizarlo. Sin red ni base de datos: funciones puras."""

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from app.models import ReadingStatus

# Claves sin tildes, en minúsculas y sin espacios sobrantes (ver `fold`).
STATUS_MAP = {
    "leido": ReadingStatus.READ,
    "no leido": ReadingStatus.PENDING,
    "no terminado": ReadingStatus.ABANDONED,
    "para consultar": ReadingStatus.REFERENCE,
    # Los valores del propio enum, para releer el CSV de revisión.
    **{status.value: status for status in ReadingStatus},
}

NO_AUTHOR = {"diversos autores", "varios autores", "aa. vv.", "aavv"}

# Partículas que van en minúscula dentro de un nombre propio.
NAME_PARTICLES = {"de", "del", "la", "las", "los", "y", "van", "von", "da", "di"}


@dataclass
class RawBook:
    import_ref: str
    title: str
    authors: list[str] = field(default_factory=list)
    status: ReadingStatus = ReadingStatus.PENDING
    notes: str | None = None


def clean(text: str | None) -> str:
    """Normaliza Unicode (NFC) y colapsa espacios."""
    if text is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(text))).strip()


def fold(text: str | None) -> str:
    """Versión para comparar: sin tildes, en minúsculas."""
    decomposed = unicodedata.normalize("NFD", clean(text).lower())
    return "".join(c for c in decomposed if unicodedata.category(c) != "Mn")


def parse_status(value: str | None) -> ReadingStatus:
    key = fold(value)
    if key not in STATUS_MAP:
        raise ValueError(f"Estado desconocido: {value!r}")
    return STATUS_MAP[key]


def fix_name_case(name: str) -> str:
    """Pone mayúscula inicial a las palabras que están enteras en minúscula.

    "Gabriel garcía Márquez" -> "Gabriel García Márquez". No toca palabras con
    mayúsculas ya escritas (McNiven, De Rojas) ni partículas (de, del...).
    Un nombre entero en mayúsculas ("DOUGLAS ADAMS") se pasa antes a minúsculas.
    """
    if name.isupper():
        name = name.lower()
    words = name.split(" ")
    fixed = [
        w.capitalize() if w.islower() and (i == 0 or w not in NAME_PARTICLES) else w
        for i, w in enumerate(words)
    ]
    return " ".join(fixed)


def parse_authors(value: str | None) -> list[str]:
    text = clean(value)
    if not text or fold(text) in NO_AUTHOR:
        return []
    return [fix_name_case(clean(part)) for part in text.split("/") if clean(part)]


def parse_row(row: tuple, source: str = "excel") -> RawBook | None:
    number, title, author, status, notes, *_ = (*row, None, None, None, None, None)
    if not clean(title):
        return None
    return RawBook(
        import_ref=f"{source}:{int(number)}",
        title=clean(title),
        authors=parse_authors(author),
        status=parse_status(status),
        notes=clean(notes) or None,
    )


def read_excel(path: Path) -> list[RawBook]:
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    books = [parse_row(row) for row in sheet.iter_rows(values_only=True)]
    return [b for b in books if b is not None]
